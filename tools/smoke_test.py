"""Exercise the published helpers with synthetic data in a temporary directory."""

import csv
import json
import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]


def run(skill, helper, *args):
    result = subprocess.run(
        [sys.executable, str(ROOT / 'skills' / skill / 'scripts' / helper),
         *map(str, args)],
        capture_output=True,
        text=True,
        encoding='utf-8',
        env={**os.environ, 'PYTHONIOENCODING': 'utf-8'},
    )
    if result.returncode:
        raise RuntimeError(f'{helper} {args[0]} failed: {result.stderr}')
    return result.stdout


def main():
    with TemporaryDirectory(prefix='office-skills-check-') as temp:
        folder = Path(temp)
        source = folder / 'report.json'
        source.write_text(json.dumps({
            'title': 'Synthetic report',
            'paragraphs': ['Chinese text: 中文测试'],
            'tables': [{'headers': ['Item', 'Count'], 'rows': [['Sample', 3]]}],
        }, ensure_ascii=False), encoding='utf-8')
        docx = folder / 'report.docx'
        run('kimi-docx', 'docx_windows.py', 'create-report', source, docx)
        document = Document(docx)
        assert document.paragraphs[0].text == 'Synthetic report'
        assert document.tables[0].cell(1, 1).text == '3'
        assert json.loads(run('kimi-docx', 'docx_windows.py', 'validate', docx))['tables'] == 1

        text = folder / 'input.txt'
        text.write_text('Synthetic PDF\nChinese text: 中文测试', encoding='utf-8')
        pdf = folder / 'report.pdf'
        run('kimi-pdf', 'pdf_windows.py', 'create-from-text', text, pdf)
        assert len(PdfReader(pdf).pages) == 1
        assert json.loads(run('kimi-pdf', 'pdf_windows.py', 'validate', pdf))['pages'] == 1
        merged = folder / 'merged.pdf'
        run('kimi-pdf', 'pdf_windows.py', 'merge', merged, pdf, pdf)
        assert len(PdfReader(merged).pages) == 2
        pages = folder / 'pages'
        run('kimi-pdf', 'pdf_windows.py', 'split', merged, pages)
        assert len(list(pages.glob('*.pdf'))) == 2
        assert 'Synthetic PDF' in run('kimi-pdf', 'pdf_windows.py', 'extract-text', merged)

        csv_path = folder / 'input.csv'
        with csv_path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerows([['部门', '数量'], ['示例', 3]])
        xlsx = folder / 'report.xlsx'
        run('kimi-xlsx', 'xlsx_windows.py', 'create-from-csv', csv_path, xlsx)
        workbook = load_workbook(xlsx)
        assert workbook['Data']['A2'].value == '示例'
        assert workbook['Data']['B2'].value == 3
        assert workbook['Data']['A1'].fill.fgColor.rgb.endswith('1F4E78')
        workbook.close()
        assert json.loads(run('kimi-xlsx', 'xlsx_windows.py', 'validate', xlsx))['rows'] == 2
    print('PASS: DOCX, PDF creation/merge/split/extraction, XLSX and validators')


if __name__ == '__main__':
    main()
