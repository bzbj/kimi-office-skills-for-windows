from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def env() -> int:
    payload = {
        "python": sys.executable,
        "pandas": pd.__version__,
        "openpyxl": openpyxl.__version__,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


def autosize(ws) -> None:
    for column in ws.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column) + 4
        ws.column_dimensions[letter].width = width
        for cell in column:
            cell.alignment = Alignment(horizontal="center", vertical="center")


def create_from_csv(input_path: Path, output_path: Path) -> int:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    dataframe = pd.read_csv(input_path, encoding="utf-8-sig")
    dataframe.to_excel(output_path, index=False, sheet_name="Data")

    workbook = load_workbook(output_path)
    sheet = workbook["Data"]
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    autosize(sheet)
    workbook.save(output_path)
    print(output_path)
    return 0


def validate(path: Path) -> int:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    payload = {
        "sheets": workbook.sheetnames,
        "rows": sheet.max_row,
        "cols": sheet.max_column,
        "a1": sheet["A1"].value,
        "a2": sheet["A2"].value,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: xlsx_windows.py {env|create-from-csv|validate} [args]")
        return 1

    command = sys.argv[1]
    if command == "env":
        return env()
    if command == "create-from-csv":
        if len(sys.argv) < 4:
            print("create-from-csv requires input.csv and output.xlsx")
            return 1
        return create_from_csv(Path(sys.argv[2]), Path(sys.argv[3]))
    if command == "validate":
        if len(sys.argv) < 3:
            print("validate requires file path")
            return 1
        return validate(Path(sys.argv[2]))

    print(f"Unknown command: {command}")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
