from __future__ import annotations

import json
import sys
from pathlib import Path

import docx
import lxml
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


def set_run_font(run, east_asia: str = "Microsoft YaHei", ascii_font: str = "Calibri", size: int = 11, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def env() -> int:
    payload = {
        "python": sys.executable,
        "python_docx_module": getattr(docx, "__file__", None),
        "lxml_version": getattr(lxml, "__version__", None),
        "default_cjk_font": "Microsoft YaHei",
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


def _add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run)


def _add_table(document: Document, table_data: dict) -> None:
    headers = table_data.get("headers", [])
    rows = table_data.get("rows", [])
    if not headers:
        return

    table = document.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        run = cell.paragraphs[0].add_run(str(text))
        set_run_font(run, bold=True)

    for row_values in rows:
        row = table.add_row().cells
        for cell, text in zip(row, row_values):
            run = cell.paragraphs[0].add_run(str(text))
            set_run_font(run)


def create_report(input_path: Path, output_path: Path) -> int:
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    title = payload.get("title")
    if title:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(str(title))
        set_run_font(run, size=18, bold=True)

    for item in payload.get("paragraphs", []):
        _add_paragraph(document, str(item))

    for table_data in payload.get("tables", []):
        _add_table(document, table_data)

    document.save(output_path)
    print(output_path)
    return 0


def validate(path: Path) -> int:
    document = Document(path)
    payload = {
        "paragraphs": len(document.paragraphs),
        "first_paragraph": document.paragraphs[0].text if document.paragraphs else "",
        "tables": len(document.tables),
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: docx_windows.py {env|create-report|validate} [args]")
        return 1

    command = sys.argv[1]
    if command == "env":
        return env()
    if command == "create-report":
        if len(sys.argv) < 4:
            print("create-report requires input.json and output.docx")
            return 1
        return create_report(Path(sys.argv[2]), Path(sys.argv[3]))
    if command == "validate":
        if len(sys.argv) < 3:
            print("validate requires file path")
            return 1
        return validate(Path(sys.argv[2]))

    print(f"Unknown command: {command}")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
